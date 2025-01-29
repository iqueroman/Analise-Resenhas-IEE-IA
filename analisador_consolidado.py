import pandas as pd
import matplotlib.pyplot as plt
from pathlib import Path
import logging
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment
import openpyxl

class AnalisadorConsolidado:
    def __init__(self, pasta_relatorios: Path):
        self.pasta_relatorios = pasta_relatorios
        self.logger = logging.getLogger('detector_ia')

    def gerar_relatorio_consolidado(self):
        """
        Gera relatório consolidado de todos os participantes
        """
        try:
            # Lista todos os arquivos Excel na pasta de relatórios
            arquivos_excel = list(self.pasta_relatorios.glob('relatório_*.xlsx'))
            
            resultados = []
            for arquivo in arquivos_excel:
                # Extrai nome do participante do nome do arquivo
                participante = arquivo.stem.replace('relatório_', '')
                
                # Lê o arquivo Excel
                df = pd.read_excel(arquivo)
                
                # Calcula as métricas existentes
                total_resenhas = len(df)
                total_gptzero_80 = len(df[df['GPTZero_Prob_IA'] >= 0.80])
                total_zerogpt_80 = len(df[df['ZeroGPT_Porcentagem_IA'] >= 80])
                total_gptzero_60 = len(df[df['GPTZero_Prob_IA'] >= 0.60])
                total_zerogpt_60 = len(df[df['ZeroGPT_Porcentagem_IA'] >= 60])
                media_gptzero = df['GPTZero_Prob_IA'].mean()
                media_zerogpt = df['ZeroGPT_Porcentagem_IA'].mean()
                
                # Novas métricas
                total_gptzero_40 = len(df[df['GPTZero_Prob_IA'] >= 0.40])
                total_zerogpt_40 = len(df[df['ZeroGPT_Porcentagem_IA'] >= 40])
                
                # Total de resenhas marcadas por qualquer detector (>=40)
                marcadas_ia = df[(df['GPTZero_Prob_IA'] >= 0.40) | 
                               (df['ZeroGPT_Porcentagem_IA'] >= 40)]
                total_marcadas_ia = len(marcadas_ia)
                
                # Percentual de resenhas marcadas
                percentual_marcadas = (total_marcadas_ia / total_resenhas * 100) if total_resenhas > 0 else 0
                
                resultados.append({
                    'Participante': participante,
                    'Total_Resenhas': total_resenhas,
                    'Total_GPTZero_80': total_gptzero_80,
                    'Total_ZeroGPT_80': total_zerogpt_80,
                    'Total_GPTZero_60': total_gptzero_60,
                    'Total_ZeroGPT_60': total_zerogpt_60,
                    'Total_GPTZero_40': total_gptzero_40,
                    'Total_ZeroGPT_40': total_zerogpt_40,
                    'Total_Marcadas_IA': total_marcadas_ia,
                    'Percentual_Marcadas_>40': percentual_marcadas,
                    'Media_GPTZero': media_gptzero,
                    'Media_ZeroGPT': media_zerogpt
                })
            
            # Cria DataFrame consolidado
            df_consolidado = pd.DataFrame(resultados)
            
            # Ordena por percentual de resenhas marcadas
            df_consolidado = df_consolidado.sort_values('Percentual_Marcadas_>40', ascending=False)
            
            # Formata a coluna de percentual para mostrar % no Excel
            df_consolidado['Percentual_Marcadas_>40'] = df_consolidado['Percentual_Marcadas_>40'].apply(lambda x: f'{x:.1f}%')
            
            # Gera arquivo Excel com formatação
            arquivo_saida = self.pasta_relatorios / 'relatorio_consolidado.xlsx'
            with pd.ExcelWriter(arquivo_saida, engine='openpyxl') as writer:
                df_consolidado.to_excel(writer, index=False, sheet_name='Consolidado')
                
                # Ajusta largura das colunas
                worksheet = writer.sheets['Consolidado']
                for idx, col in enumerate(df_consolidado.columns):
                    # Encontra o comprimento máximo na coluna
                    max_length = max(
                        df_consolidado[col].astype(str).apply(len).max(),  # Maior valor
                        len(str(col))  # Tamanho do cabeçalho
                    )
                    # Ajusta a largura (um pouco maior para garantir a legibilidade)
                    worksheet.column_dimensions[get_column_letter(idx + 1)].width = max_length + 2

                # Formata cabeçalhos
                for cell in worksheet[1]:
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal='center', wrap_text=True)

                # Adiciona comentários explicativos
                explicacoes = {
                    'Participante': 'Nome do participante extraído do nome do arquivo',
                    'Total_Resenhas': 'Número total de resenhas enviadas pelo participante',
                    'Total_GPTZero_80': 'Número de resenhas com GPTZero_Prob_IA >= 0.80',
                    'Total_ZeroGPT_80': 'Número de resenhas com ZeroGPT_Porcentagem_IA >= 80',
                    'Total_GPTZero_60': 'Número de resenhas com GPTZero_Prob_IA >= 0.60',
                    'Total_ZeroGPT_60': 'Número de resenhas com ZeroGPT_Porcentagem_IA >= 60',
                    'Total_GPTZero_40': 'Número de resenhas com GPTZero_Prob_IA >= 0.40',
                    'Total_ZeroGPT_40': 'Número de resenhas com ZeroGPT_Porcentagem_IA >= 40',
                    'Total_Marcadas_IA': 'Número de resenhas que têm (GPTZero_Prob_IA >= 0.40 OU ZeroGPT_Porcentagem_IA >= 40)',
                    'Percentual_Marcadas_>40': 'Total_Marcadas_IA dividido pelo Total_Resenhas, multiplicado por 100',
                    'Media_GPTZero': 'Média do GPTZero_Prob_IA de todas as resenhas do participante',
                    'Media_ZeroGPT': 'Média do ZeroGPT_Porcentagem_IA de todas as resenhas do participante'
                }

                # Adiciona os comentários nas células
                for idx, col in enumerate(df_consolidado.columns):
                    if col in explicacoes:
                        cell = worksheet[f"{get_column_letter(idx + 1)}1"]
                        cell.comment = openpyxl.comments.Comment(
                            explicacoes[col],
                            'Detector IA'
                        )
            
            self.logger.info(f"Relatório consolidado gerado: {arquivo_saida}")
            return arquivo_saida
            
        except Exception as e:
            self.logger.error(f"Erro ao gerar relatório consolidado: {str(e)}", exc_info=True)
            raise 